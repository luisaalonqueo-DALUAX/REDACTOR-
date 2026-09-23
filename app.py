import streamlit as st
import sqlite3
import re
import io
from datetime import datetime

# Intentar cargar la librería para Word
try:
    import docx
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_DISPONIBLE = True
except ImportError:
    DOCX_DISPONIBLE = False

st.set_page_config(
    page_title="Redactor de Documentos Oficiales - Chubut",
    page_icon="📝",
    layout="wide"
)

# ---------------------------------------------------------
# FUNCIONES DE FORMATO Y LIMPIEZA
# ---------------------------------------------------------
MESES_ESPANOL = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril",
    5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
    9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"
}

def obtener_fecha_actual_espanol():
    hoy = datetime.now()
    mes = MESES_ESPANOL[hoy.month]
    return f"Rawson, {hoy.day} de {mes} de {hoy.year}.-"

def limpiar_texto(texto):
    if not texto:
        return ""
    texto_limpio = re.sub(r'[ \t]+', ' ', texto)
    return texto_limpio.strip()

# ---------------------------------------------------------
# GENERACIÓN DE ARCHIVO WORD (.DOCX) EDITABLE
# ---------------------------------------------------------
def generar_documento_word(tipo_doc, lugar_fecha, dest_final, asunto_ref, cuerpo, pie_formateado, es_duplicado):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    def agregar_bloque(doc_obj):
        p_leyenda = doc_obj.add_paragraph()
        p_leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_ley = p_leyenda.add_run('“Año de la Innovación y Modernización del Estado de la Provincia del Chubut”')
        run_ley.italic = True
        run_ley.font.name = 'Verdana'
        run_ley.font.size = Pt(9.5)
        run_ley.font.bold = True
        run_ley.font.color.rgb = RGBColor(45, 55, 72)
        
        p_enc = doc_obj.add_paragraph()
        p_enc.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_enc1 = p_enc.add_run('Gobierno del Chubut | Ministerio de Producción\n')
        run_enc1.bold = True
        run_enc1.font.name = 'Verdana'
        run_enc1.font.size = Pt(11)
        
        run_enc2 = p_enc.add_run('Subsecretaría de Financiamiento y Comercio para la Producción\nDepartamento de Administración de Personal')
        run_enc2.font.name = 'Verdana'
        run_enc2.font.size = Pt(9)
        run_enc2.font.color.rgb = RGBColor(74, 85, 104)
        
        p_fecha = doc_obj.add_paragraph()
        p_fecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_fecha = p_fecha.add_run(lugar_fecha)
        run_fecha.bold = True
        run_fecha.font.name = 'Verdana'
        run_fecha.font.size = Pt(11)
        
        p_dest = doc_obj.add_paragraph()
        p_dest.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_dest = p_dest.add_run(dest_final)
        run_dest.font.name = 'Verdana'
        run_dest.font.size = Pt(11)
        p_dest.paragraph_format.line_spacing = 1.3
        
        if asunto_ref:
            p_ref = doc_obj.add_paragraph()
            p_ref.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run_ref_b = p_ref.add_run('Ref.: ')
            run_ref_b.bold = True
            run_ref_b.font.name = 'Verdana'
            run_ref_b.font.size = Pt(11)
            run_ref_t = p_ref.add_run(asunto_ref)
            run_ref_t.font.name = 'Verdana'
            run_ref_t.font.size = Pt(11)
            
        p_cuerpo = doc_obj.add_paragraph()
        p_cuerpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        if tipo_doc == "Nota Oficial":
            texto_cuerpo = f"De mi mayor consideración:\n\n{cuerpo}\n\nSin otro particular, saludo a Ud. atentamente."
        elif tipo_doc == "Memorándum":
            texto_cuerpo = f"{cuerpo}\n\nSin más que agregar, saludo a usted muy atentamente."
        else:
            texto_cuerpo = cuerpo

        run_cuerpo = p_cuerpo.add_run(texto_cuerpo)
        run_cuerpo.font.name = 'Verdana'
        run_cuerpo.font.size = Pt(11)
        p_cuerpo.paragraph_format.line_spacing = 1.5
        
        p_pie = doc_obj.add_paragraph()
        p_pie.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_pie = p_pie.add_run(f"\n{pie_formateado}")
        run_pie.bold = True
        run_pie.font.name = 'Verdana'
        run_pie.font.size = Pt(10)
        
        p_cont = doc_obj.add_paragraph()
        p_cont.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_cont = p_cont.add_run("Mariano Moreno y Luis Costa | Rawson | Chubut (Teléfono 280-4485125/126)")
        run_cont.font.name = 'Verdana'
        run_cont.font.size = Pt(8)
        run_cont.font.color.rgb = RGBColor(128, 128, 128)

    agregar_bloque(doc)
    
    if es_duplicado:
        p_corte = doc.add_paragraph()
        p_corte.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_corte = p_corte.add_run("\n- - - - - - - - - - - - CORTAR AQUÍ (ORIGINAL Y DUPLICADO) - - - - - - - - - - - -\n")
        run_corte.font.name = 'Verdana'
        run_corte.font.size = Pt(8)
        run_corte.font.color.rgb = RGBColor(160, 174, 192)
        
        agregar_bloque(doc)
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ---------------------------------------------------------
# BASE DE DATOS LOCAL
# ---------------------------------------------------------
def init_db():
    conn = sqlite3.connect("documentos.db")
    c = conn.cursor()
    c.execute('''
        CREATE TABLE IF NOT EXISTS documentos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tipo TEXT,
            numero INTEGER,
            anio TEXT,
            fecha TEXT,
            destinatario TEXT,
            asunto TEXT,
            cuerpo TEXT,
            iniciales TEXT
        )
    ''')
    c.execute('''
        CREATE TABLE IF NOT EXISTS destinatarios_frecuentes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT UNIQUE
        )
    ''')
    
    frecuentes = [
        "Departamento Sueldos",
        "Mesa General de Entradas y Salidas\nMinisterio de Producción\nSU DESPACHO",
        "Ministerio de Producción\nDirección General de Asuntos Legales\nAbg. Norma Navarro\nSU DESPACHO:",
        "División Sueldos y Cargas Sociales",
        "Subsecretaría de Financiamiento y Comercio para la Producción"
    ]
    for dest in frecuentes:
        c.execute('INSERT OR IGNORE INTO destinatarios_frecuentes (nombre) VALUES (?)', (dest,))
        
    conn.commit()
    conn.close()

def obtener_destinatarios_frecuentes():
    conn = sqlite3.connect("documentos.db")
    c = conn.cursor()
    c.execute('SELECT nombre FROM destinatarios_frecuentes ORDER BY nombre ASC')
    filas = c.fetchall()
    conn.close()
    return [f[0] for f in filas]

def guardar_destinatario_nuevo(nombre):
    if nombre and len(nombre.strip()) > 3:
        conn = sqlite3.connect("documentos.db")
        c = conn.cursor()
        c.execute('INSERT OR IGNORE INTO destinatarios_frecuentes (nombre) VALUES (?)', (limpiar_texto(nombre),))
        conn.commit()
        conn.close()

def obtener_ultimo_numero(tipo, anio):
    conn = sqlite3.connect("documentos.db")
    c = conn.cursor()
    c.execute("SELECT MAX(numero) FROM documentos WHERE tipo = ? AND anio = ?", (tipo, anio))
    res = c.fetchone()[0]
    conn.close()
    return res if res else 0

def guardar_documento(tipo, numero, anio, fecha, destinatario, asunto, cuerpo, iniciales):
    conn = sqlite3.connect("documentos.db")
    c = conn.cursor()
    c.execute('''
        INSERT INTO documentos (tipo, numero, anio, fecha, destinatario, asunto, cuerpo, iniciales)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    ''', (tipo, numero, anio, fecha, destinatario, asunto, cuerpo, iniciales))
    conn.commit()
    conn.close()
    guardar_destinatario_nuevo(destinatario)

def obtener_todos_documentos():
    conn = sqlite3.connect("documentos.db")
    c = conn.cursor()
    c.execute('SELECT id, tipo, numero, anio, fecha, destinatario, asunto FROM documentos ORDER BY id DESC')
    registros = c.fetchall()
    conn.close()
    return registros

def buscar_documentos(query=""):
    conn = sqlite3.connect("documentos.db")
    c = conn.cursor()
    if query:
        c.execute('''
            SELECT id, tipo, numero, anio, fecha, destinatario, asunto 
            FROM documentos 
            WHERE destinatario LIKE ? OR asunto LIKE ? OR cuerpo LIKE ? OR tipo LIKE ?
            ORDER BY id DESC
        ''', (f'%{query}%', f'%{query}%', f'%{query}%', f'%{query}%'))
    else:
        c.execute('SELECT id, tipo, numero, anio, fecha, destinatario, asunto FROM documentos ORDER BY id DESC')
    registros = c.fetchall()
    conn.close()
    return registros

init_db()

# ---------------------------------------------------------
# INTERFAZ
# ---------------------------------------------------------
st.title("📝 Redactor de Documentos Oficiales")
st.subheader("Gobierno de la Provincia del Chubut - Ministerio de Producción")

pestana1, pestana2 = st.tabs(["✍️ Redactar Documento", "🔍 Buscador e Índice General"])

with pestana1:
    col_redaccion, col_indice = st.columns([2, 1])
    
    with col_indice:
        st.markdown("### 📋 Índice de Correlatividad")
        st.caption("Historial de números utilizados")
        
        docs_registrados = obtener_todos_documentos()
        
        if docs_registrados:
            for doc in docs_registrados:
                d_id, d_tipo, d_num, d_anio, d_fecha, d_dest, d_asunto = doc
                st.markdown(f"**{d_tipo} N° {d_num}/{d_anio}**")
                st.markdown(f"• **Fecha:** {d_fecha}")
                st.markdown(f"• **Motivo/Ref:** {d_asunto}")
                st.markdown(f"• **Destino:** {d_dest[:40]}...")
                st.markdown("---")
        else:
            st.info("Aún no hay documentos registrados.")

    with col_redaccion:
        anio_actual = str(datetime.now().year)
        
        st.markdown("### 1. Destinatario")
        
        opciones_destinatarios = ["-- Escribir un nuevo destinatario --"] + obtener_destinatarios_frecuentes()
        destinatario_seleccionado = st.selectbox("Seleccionar un destinatario habitual:", opciones_destinatarios)
        
        if destinatario_seleccionado == "-- Escribir un nuevo destinatario --":
            destinatario_input = st.text_area("Ingresa Destinatario:", value="Dirección General de Asuntos Legales\nAbg. Norma Navarro\nSU DESPACHO:")
        else:
            destinatario_input = destinatario_seleccionado

        destinatario_input = limpiar_texto(destinatario_input)
        
        sugerencia = "Nota Oficial"
        explicacion = "Las comunicaciones dirigidas a Asuntos Legales u otros Ministerios/Direcciones externas corresponden a Nota Oficial."
        
        dest_lower = destinatario_input.lower()
        if any(p in dest_lower for p in ["sueldos", "división", "departamento sueldos", "liquidacion", "costos"]):
            sugerencia = "Memorándum"
            explicacion = "Las comunicaciones internas a dependencias propias corresponden a Memorándum."
        elif any(p in dest_lower for p in ["mesa", "entradas", "acumular", "salidas"]):
            sugerencia = "Pase Administrativo"
            explicacion = "El giro o acumulación de expedientes corresponde a Pase Administrativo."
        
        if destinatario_input:
            st.info(f"💡 **Sugerencia Normativa:** Según el destino, corresponde un/a **{sugerencia}**.\n_{explicacion}_")

        tipo_doc = st.selectbox(
            "Tipo de Documento a emitir:",
            ["Nota Oficial", "Memorándum", "Pase Administrativo"],
            index=["Nota Oficial", "Memorándum", "Pase Administrativo"].index(sugerencia)
        )

        ultimo_num = obtener_ultimo_numero(tipo_doc, anio_actual)
        siguiente_num = ultimo_num + 1

        st.warning(f"📌 Próximo número disponible para **{tipo_doc}**: **{siguiente_num}**.")

        with st.form("form_redaccion"):
            c1, c2 = st.columns(2)
            with c1:
                num_doc = st.number_input("Número Correlativo:", value=siguiente_num, step=1)
                anio_doc = st.text_input("Año:", value=anio_actual)
            with c2:
                lugar_fecha = st.text_input("Lugar y Fecha:", value=obtener_fecha_actual_espanol())
                iniciales = st.text_input("Iniciales al pie:", value="L.I.A.")

            st.markdown("---")
            
            if tipo_doc == "Memorándum":
                de_persona = st.text_input("Producido por:", value="Departamento de Administración de Personal")
                para_persona = st.text_input("Para Información:", value="Departamento Sueldos")
                asunto_ref = st.text_input("Ref. / Expte.:", value="Expte: N° 1729/2026-MP- Reclamo Adicional...")
                cuerpo = st.text_area("Cuerpo del Memorándum:", height=180)
                dest_final = f"Producido por: {limpiar_texto(de_persona)}\nPara Información: {limpiar_texto(para_persona)}"
            elif tipo_doc == "Pase Administrativo":
                dest_final = destinatario_input if destinatario_input else "Mesa General de Entradas y Salidas\nMinisterio de Producción\nSU DESPACHO"
                st.text_area("Destinatario:", value=dest_final)
                asunto_ref = st.text_input("Ref. / Expte. a acumular:", value="Acumulación de Expediente N° 3091/2023...")
                cuerpo = st.text_area("Cuerpo del Pase:", height=180, value="Por medio del presente me dirijo a usted, con el fin de solicitar acumular Expediente Nº 3091/2023 – MAGIyC – constan 14 fojas, al Expediente 1612/2024 el mismo consta con 30 fojas cuyo extracto exprese lo siguiente:\n\nS/ Recaratular Expte. N°3498/2023 UEP.MAGYC-Ref. Pase N° 471/2023 SAP-UEP-MAGIyC “Baja por fallecimiento del agente OLATE, Juan Carlos”.\n\nUna vez cumplido vuelva a este Departamento de Personal, para proseguir con el trámite administrativo correspondiente.")
            else: # Nota Oficial
                dest_final = destinatario_input
                st.text_area("Destinatario completo:", value=dest_final)
                asunto_ref = st.text_input("Ref. / Expte.:")
                cuerpo = st.text_area("Cuerpo de la Nota:", height=180)

            asunto_ref = limpiar_texto(asunto_ref)
            cuerpo = limpiar_texto(cuerpo)
            
            submitted = st.form_submit_button("🚀 Generar Documento y Guardar en el Índice")

        if submitted:
            guardar_documento(tipo_doc, num_doc, anio_doc, lugar_fecha, dest_final, asunto_ref, cuerpo, iniciales)
            st.success(f"✅ Registrado: **{tipo_doc} N° {num_doc}/{anio_doc}**.")
            st.rerun()

        leyenda_oficial = "“Año de la Innovación y Modernización del Estado de la Provincia del Chubut”"
        contacto_pie = "Mariano Moreno y Luis Costa | Rawson | Chubut (Teléfono 280-4485125/126)"
        
        if tipo_doc == "Memorándum":
            pie_formateado = f"Memo Nº {siguiente_num}/{anio_actual} –DAP- SsFyCP-MP ({iniciales})"
            doc_cuerpo_completo = f"{cuerpo}\n\nSin más que agregar, saludo a usted muy atentamente.\n\n{pie_formateado}"
            es_duplicado = False
        elif tipo_doc == "Pase Administrativo":
            pie_formateado = f"Pase N° {siguiente_num}/{anio_actual}-dap-SsFyCP-MP({iniciales})"
            doc_cuerpo_completo = f"{cuerpo}\n\n{pie_formateado}"
            es_duplicado = True
        else: # Nota Oficial
            pie_formateado = f"Nota N° {siguiente_num}/{anio_actual} DAP-SsFyCP- Ministerio de Producción ({iniciales})"
            doc_cuerpo_completo = f"De mi mayor consideración:\n\n{cuerpo}\n\nSin otro particular, saludo a Ud. atentamente.\n\n({iniciales}) {pie_formateado}"
            es_duplicado = False

        st.markdown("---")
        st.subheader("📄 Vista Previa Texto")
        
        # VISTA PREVIA LIMPIA NATIVA EN STREAMLIT
        with st.container(border=True):
            st.caption(leyenda_oficial)
            st.markdown("**Gobierno del Chubut | Ministerio de Producción**")
            st.text("Subsecretaría de Financiamiento y Comercio para la Producción\nDepartamento de Administración de Personal")
            st.markdown(f"<div style='text-align: right;'><b>{lugar_fecha}</b></div>", unsafe_allow_html=True)
            st.text(dest_final)
            if asunto_ref:
                st.markdown(f"<div style='text-align: right;'><b>Ref.:</b> {asunto_ref}</div>", unsafe_allow_html=True)
            st.write(doc_cuerpo_completo)
            if es_duplicado:
                st.divider()
                st.caption("✂️ (EN PASE ADMINISTRATIVO SE IMPRIME EL DUPLICADO ABAJO EN LA MISMA HOJA)")
            st.caption(f"📍 {contacto_pie}")

        st.markdown("---")
        col_btn1, col_btn2 = st.columns(2)
        
        if DOCX_DISPONIBLE:
            file_word = generar_documento_word(tipo_doc, lugar_fecha, dest_final, asunto_ref, cuerpo, pie_formateado, es_duplicado)
            with col_btn1:
                st.download_button(
                    label="📄 Descargar en Word (.docx) Editable",
                    data=file_word,
                    file_name=f"{tipo_doc.lower().replace(' ', '_')}_{siguiente_num}_{anio_actual}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
        else:
            with col_btn1:
                st.warning("Agregar 'python-docx' a requirements.txt para Word")

        doc_texto_descarga = f"{lugar_fecha}\n\n{dest_final}\nRef.: {asunto_ref}\n\n{doc_cuerpo_completo}"
        with col_btn2:
            st.download_button(
                label="📥 Descargar Texto (.txt)",
                data=doc_texto_descarga,
                file_name=f"{tipo_doc.lower().replace(' ', '_')}_{siguiente_num}_{anio_actual}.txt",
                mime="text/plain"
            )

with pestana2:
    st.markdown("### 🔍 Buscador General de Documentos")
    busqueda = st.text_input("Buscar por expediente, motivo, asunto, destinatario o tipo:", value="")
    
    resultados = buscar_documentos(busqueda)
    
    if resultados:
        for item in resultados:
            doc_id, t_doc, n_doc, a_doc, f_doc, d_doc, as_doc = item
            with st.expander(f"📌 {t_doc} N° {n_doc}/{a_doc} — Ref: {as_doc}"):
                st.write(f"**Fecha:** {f_doc}")
                st.write(f"**Destinatario:** {d_doc}")
                st.write(f"**Referencia / Motivo:** {as_doc}")
    else:
        st.info("No se encontraron documentos registrados con ese criterio de búsqueda.")
